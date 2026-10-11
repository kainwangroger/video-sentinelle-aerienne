"""Convertit la fiche technique Markdown en document Word (.docx), schémas inclus.

Gère : titres, paragraphes, gras/italique/code en ligne, listes, tableaux, citations,
images locales et blocs de code (le diagramme Mermaid est ignoré : le schéma PNG le suit).

Usage : python3 drone/outils/md_vers_docx.py [entrée.md] [sortie.docx]
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLEU = RGBColor(0x0B, 0x3D, 0x6E)
CYAN = RGBColor(0x00, 0x7F, 0xA8)


def _fond_cellule(cellule, couleur):
    tc = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tc.append(shd)


def ajouter_texte(paragraphe, texte, gras_par_defaut=False):
    """Ajoute du texte avec **gras**, *italique* et `code`, sans les liens Markdown."""
    texte = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texte)
    morceaux = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texte)
    for m in morceaux:
        if not m:
            continue
        if m.startswith("**") and m.endswith("**"):
            run = paragraphe.add_run(m[2:-2])
            run.bold = True
        elif m.startswith("*") and m.endswith("*") and len(m) > 2:
            run = paragraphe.add_run(m[1:-1])
            run.italic = True
        elif m.startswith("`") and m.endswith("`"):
            run = paragraphe.add_run(m[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = CYAN
        else:
            run = paragraphe.add_run(m)
            run.bold = gras_par_defaut or None


def convertir(entree, sortie):
    entree = Path(entree)
    lignes = entree.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for niveau in range(1, 4):
        doc.styles[f"Heading {niveau}"].font.color.rgb = BLEU

    i = 0
    dans_sommaire = False
    while i < len(lignes):
        ligne = lignes[i]
        brut = ligne.strip()
        if brut.startswith("```"):
            langage = brut[3:].strip()
            bloc = []
            i += 1
            while i < len(lignes) and not lignes[i].strip().startswith("```"):
                bloc.append(lignes[i])
                i += 1
            i += 1
            if langage == "mermaid":
                pass  # le schéma PNG équivalent suit le bloc dans le Markdown
            else:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(bloc))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            continue
        if brut.startswith("#"):
            niveau = len(brut) - len(brut.lstrip("#"))
            titre = brut[niveau:].strip()
            dans_sommaire = titre.lower() == "sommaire"
            if niveau == 1:
                p = doc.add_heading(titre, 0)
            else:
                doc.add_heading(titre, min(niveau - 1, 3))
            i += 1
            continue
        if dans_sommaire and re.match(r"^\d+\. \[", brut):
            i += 1
            continue
        if brut.startswith("![") and "](" in brut:
            chemin = re.search(r"\(([^)]+)\)", brut).group(1)
            image = entree.parent / chemin
            if image.exists():
                doc.add_picture(str(image), width=Cm(16))
                legende = re.search(r"!\[([^\]]*)\]", brut).group(1)
                p = doc.add_paragraph(legende)
                p.runs[0].italic = True
                p.runs[0].font.size = Pt(9)
            i += 1
            continue
        if brut.startswith("|"):
            tableau = []
            while i < len(lignes) and lignes[i].strip().startswith("|"):
                cellules = [c.strip() for c in lignes[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cellules):
                    tableau.append(cellules)
                i += 1
            n = max(len(r) for r in tableau)
            t = doc.add_table(rows=len(tableau), cols=n)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, rangee in enumerate(tableau):
                for c in range(n):
                    cellule = t.cell(r, c)
                    cellule.text = ""
                    ajouter_texte(cellule.paragraphs[0], rangee[c] if c < len(rangee) else "", gras_par_defaut=(r == 0))
                    for run in cellule.paragraphs[0].runs:
                        run.font.size = Pt(9.5)
                    if r == 0:
                        _fond_cellule(cellule, "D9EEF5")
            doc.add_paragraph()
            continue
        if brut.startswith(">"):
            texte = []
            while i < len(lignes) and lignes[i].strip().startswith(">"):
                texte.append(lignes[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            ajouter_texte(p, " ".join(texte))
            for run in p.runs:
                run.font.color.rgb = BLEU
            continue
        m_liste = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ligne)
        if m_liste:
            retrait, puce, texte = m_liste.groups()
            texte = re.sub(r"^\[[ x]\]\s*", "☐ ", texte)
            style_liste = "List Number" if puce[0].isdigit() else "List Bullet"
            if len(retrait) >= 2:
                style_liste += " 2"
            p = doc.add_paragraph(style=style_liste)
            # rattacher les lignes de continuation
            i += 1
            while i < len(lignes) and lignes[i].startswith("  ") and lignes[i].strip() and \
                    not re.match(r"^\s*([-*]|\d+\.)\s+", lignes[i]):
                texte += " " + lignes[i].strip()
                i += 1
            ajouter_texte(p, texte)
            continue
        if brut in ("---", ""):
            i += 1
            continue
        texte = [brut]
        i += 1
        while i < len(lignes) and lignes[i].strip() and not re.match(r"^(#|\||>|!\[|```|\s*([-*]|\d+\.)\s)", lignes[i]):
            texte.append(lignes[i].strip())
            i += 1
        ajouter_texte(doc.add_paragraph(), " ".join(texte))
    doc.save(sortie)


if __name__ == "__main__":
    racine = Path(__file__).resolve().parents[1]
    entree = Path(sys.argv[1]) if len(sys.argv) > 1 else racine / "fiche_technique_complete.md"
    sortie = Path(sys.argv[2]) if len(sys.argv) > 2 else racine / "fiche_technique_complete.docx"
    convertir(entree, sortie)
    print(f"Écrit : {sortie}")
