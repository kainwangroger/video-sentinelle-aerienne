"""Extrait le contenu de fiche_technique.docx en Markdown (output/fiche_extraite.md).

Sert de référence pour écrire la narration de src/content.py :
paragraphes avec leurs niveaux de titre, puis tableaux.
"""

import sys
from pathlib import Path

import docx

RACINE = Path(__file__).resolve().parent.parent


def vers_markdown(chemin):
    doc = docx.Document(chemin)
    lignes = []
    for p in doc.paragraphs:
        s = p.text.strip()
        if not s:
            continue
        style = p.style.name
        if style == "Title":
            lignes.append(f"# {s}")
        elif style.startswith("Heading"):
            niveau = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
            lignes.append(f"\n{'#' * (niveau + 1)} {s}\n")
        elif "List" in style:
            lignes.append(f"- {s}")
        else:
            lignes.append(f"\n{s}\n")
    for i, table in enumerate(doc.tables, 1):
        lignes.append(f"\n### Tableau {i}\n")
        for j, rangee in enumerate(table.rows):
            cellules = [c.text.strip().replace("\n", " ").replace("|", "/") for c in rangee.cells]
            lignes.append("| " + " | ".join(cellules) + " |")
            if j == 0:
                lignes.append("|" + "---|" * len(cellules))
    return "\n".join(lignes) + "\n"


if __name__ == "__main__":
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else RACINE / "fiche_technique.docx"
    cible = RACINE / "output" / "fiche_extraite.md"
    cible.parent.mkdir(exist_ok=True)
    cible.write_text(vers_markdown(source), encoding="utf-8")
    print(f"Extrait : {cible}")
